"""Build editable Chinese manuscript with native Office Math; no result generation."""
from pathlib import Path
import re,json,zipfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'paper/manuscript_source.md'
OUT=ROOT/'paper/TAAS-FPO_论文稿_待填实验结果.docx'
def el(tag, **attrs):
    x=OxmlElement(tag)
    for k,v in attrs.items(): x.set(qn(k),str(v))
    return x
def mr(text):
    x=el('m:r');t=el('m:t');t.set(qn('xml:space'),'preserve');t.text=text;x.append(t);return x
def box(tag, nodes):
    x=el(tag)
    for n in nodes:x.append(n)
    return x

class Math:
    """Small reviewed subset: groups, fractions, roots, native sub/superscripts."""
    def __init__(self,s):
        self.s=s.replace(r'\left','').replace(r'\right','').replace(r'\qquad','   ').replace(r'\quad','  ')
        self.s=self.s.replace(r'\max','max').replace(r'\exp','exp').replace(r'\bar r',r'\bar{r}').replace('Σ','∑')
        self.i=0
    def atom(self):
        s=self.s
        if self.i>=len(s):return []
        if s[self.i]=='{':
            self.i+=1;n=self.parse('}');self.i+=1;return n
        for cmd,accent in [(r'\dot','̇'),(r'\bar','̅')]:
            if s.startswith(cmd,self.i):
                self.i+=len(cmd);arg=self.atom();x=el('m:acc')
                pr=el('m:accPr');pr.append(el('m:chr',**{'m:val':accent}))
                x.append(pr);x.append(box('m:e',arg));return [x]
        if s.startswith(r'\frac',self.i):
            self.i+=5;num=self.atom();den=self.atom()
            f=el('m:f');f.append(box('m:num',num));f.append(box('m:den',den));return [f]
        if s.startswith(r'\sqrt',self.i):
            self.i+=5;arg=self.atom();x=el('m:rad');pr=el('m:radPr');pr.append(el('m:degHide',**{'m:val':'1'}))
            x.append(pr);x.append(el('m:deg'));x.append(box('m:e',arg));return [x]
        c=s[self.i];self.i+=1
        if c=='\\':raise ValueError('Unsupported math command '+s[self.i:])
        return [mr(c)]
    def parse(self,end=None):
        nodes=[]
        while self.i<len(self.s) and self.s[self.i]!=end:
            base=self.atom();sub=sup=None
            while self.i<len(self.s) and self.s[self.i] in '_^':
                typ=self.s[self.i];self.i+=1
                v=self.atom()
                if typ=='_':sub=v
                else:sup=v
            if sub is not None or sup is not None:
                tag='m:sSubSup' if sub is not None and sup is not None else ('m:sSub' if sub is not None else 'm:sSup')
                x=el(tag);x.append(box('m:e',base))
                if sub is not None:x.append(box('m:sub',sub))
                if sup is not None:x.append(box('m:sup',sup))
                nodes.append(x)
            else:nodes.extend(base)
        return nodes

def font(run,size=11,bold=False,east='宋体'):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    rp=run._element.get_or_add_rPr();fonts=rp.rFonts
    if fonts is None:fonts=el('w:rFonts');rp.insert(0,fonts)
    fonts.set(qn('w:eastAsia'),east)
def para(text,style=None):
    p=doc.add_paragraph(style=style)
    pattern=r'(?:[A-Za-z]+|[Α-Ωα-ω])_(?:[A-Za-z0-9]+|[Α-Ωα-ω])'
    parts=re.split('('+pattern+')',text)
    for part in parts:
        if re.fullmatch(pattern,part):
            base,sub=part.split('_');om=el('m:oMath')
            for node in Math(base+'_{'+sub+'}').parse():om.append(node)
            p._p.append(om)
        else:font(p.add_run(part))
    p.paragraph_format.widow_control=True
    return p
def equation(s,number):
    # Two-column borderless table preserves right aligned numbering.
    t=doc.add_table(rows=1,cols=2);t.autofit=False
    t.columns[0].width=Inches(5.98);t.columns[1].width=Inches(.52)
    for row in t.rows:row._tr.get_or_add_trPr().append(el('w:cantSplit'))
    p=t.cell(0,0).paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.line_spacing=1.0
    om=el('m:oMath');nodes=Math(s).parse()
    for n in nodes:om.append(n)
    p._p.append(om)
    np=t.cell(0,1).paragraphs[0];np.alignment=WD_ALIGN_PARAGRAPH.RIGHT;font(np.add_run('('+str(number)+')'),10)
    for c in t.rows[0].cells:
        tcpr=c._tc.get_or_add_tcPr();m=el('w:tcMar')
        for edge in ['top','bottom','left','right']:m.append(el('w:'+edge,**{'w:w':'20','w:type':'dxa'}))
        tcpr.append(m)
    return t
def table(caption,lines):
    p=para(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    for x in p.runs:font(x,10.5,True)
    rows=[[x.strip() for x in row.split('|')] for row in lines];cols=len(rows[0])
    t=doc.add_table(rows=0,cols=cols);t.autofit=False
    widths={3:[1.1,2.6,2.8],4:[1.5,1.8,1.6,1.6],5:[1.5,1.35,1.35,1.15,1.15],6:[.55,.55,.55,1.4,1.6,1.85]}[cols]
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    borders=el('w:tblBorders')
    for side in ['top','bottom']:borders.append(el('w:'+side,**{'w:val':'single','w:sz':'8','w:color':'000000'}))
    for side in ['left','right','insideH','insideV']:borders.append(el('w:'+side,**{'w:val':'nil'}))
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells;t.rows[-1]._tr.get_or_add_trPr().append(el('w:cantSplit'))
        if ri==0:t.rows[-1]._tr.get_or_add_trPr().append(el('w:tblHeader'))
        for ci,txt in enumerate(row):
            cells[ci].width=Inches(widths[ci])
            p=cells[ci].paragraphs[0];p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
            font(p.add_run(txt),9.5,ri==0)
            if ri==0:
                b=el('w:tcBorders');b.append(el('w:bottom',**{'w:val':'single','w:sz':'4'}));cells[ci]._tc.get_or_add_tcPr().append(b)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

doc=Document()
sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.35)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
normal.element.get_or_add_rPr().append(el('w:rFonts',**{'w:eastAsia':'宋体'}))
normal.paragraph_format.line_spacing=1.2;normal.paragraph_format.space_after=Pt(5)
for level,size in [(1,14),(2,12),(3,11)]:
    st=doc.styles['Heading '+str(level)];st.font.color.rgb=RGBColor(0,0,0);st.font.name='Times New Roman';st.font.size=Pt(size)
    st.element.get_or_add_rPr().append(el('w:rFonts',**{'w:eastAsia':'黑体'}))
    st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6)
header=sec.header.paragraphs[0];font(header.add_run('TAAS-FPO  ·  方法与实验协议稿  ·  实验结果待填'),8.5);header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run('— '),9)
r=footer.add_run();r._r.append(el('w:fldChar',**{'w:fldCharType':'begin'}));it=el('w:instrText');it.text=' PAGE ';r._r.append(it);r._r.append(el('w:fldChar',**{'w:fldCharType':'end'}))
font(footer.add_run(' —'),9)
doc.core_properties.title='面向动态山地无人机节能路径规划的地形感知与成本感知自适应游隼优化方法'
doc.core_properties.subject='方法与实验协议；真实结果待填'
doc.core_properties.author='待作者填写';doc.core_properties.keywords='TAAS-FPO;UAV;pending experimental results'
lines=SOURCE.read_text(encoding='utf-8').splitlines();i=0;n=0;nt=0
while i<len(lines):
    text=lines[i].strip();i+=1
    if not text:continue
    if text.startswith('@eq '):n+=1;equation(text[4:],n)
    elif text.startswith('@table '):
        cap=text[7:];rows=[]
        while i<len(lines) and lines[i].strip()!='@end':rows.append(lines[i]);i+=1
        i+=1;table(cap,rows);nt+=1
    elif text.startswith('# '):
        p=para(text[2:].replace('规划的地形','规划的\n地形'));p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:font(run,18,True,'黑体')
    elif text.startswith('## '):
        title=text[3:]
        if title in ['1 引言','参考文献']:doc.add_page_break()
        p=doc.add_paragraph(title,style='Heading 1')
    elif text.startswith('### '):doc.add_paragraph(text[4:],style='Heading 2')
    else:
        p=para(text)
        if text.startswith('['):p.paragraph_format.first_line_indent=Pt(0)
        elif text.startswith(('作者','稿件状态','关键词','Keywords')):p.paragraph_format.first_line_indent=Pt(0)
        else:p.paragraph_format.first_line_indent=Pt(22)
        if text.startswith('稿件状态'):
            for run in p.runs:font(run,9)
OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
with zipfile.ZipFile(OUT) as z:
    xml=z.read('word/document.xml')
    assert xml.count(b'<m:oMath>')>=n and b'<m:f>' in xml and b'<m:rad>' in xml
(ROOT/'paper/document_build_check.json').write_text(json.dumps({'equations':n,'tables':nt,'native_math':True,'results_filled':False,'submission_ready':False,'file':OUT.name},ensure_ascii=False,indent=2),encoding='utf-8')
print(OUT);print('Native editable equations:',n,'Tables:',nt)
